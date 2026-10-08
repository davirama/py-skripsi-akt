import pytest
from docx import Document

from core.paths import (
    pilih_template_berdasarkan_pembimbing,
    pilih_template_nota_dinas_berdasarkan_pembimbing,
)
from core.word_generator import generate_docx, sanitize_filename


def teks_docx(path) -> str:
    d = Document(str(path))
    bagian = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                bagian.append(c.text)
    return "\n".join(bagian)


@pytest.mark.parametrize(
    "raw,hasil",
    [
        ('a/b\\c:d*e?f"g<h>i|j', "abcdefghij"),
        ("  Andi    Wijaya  ", "Andi Wijaya"),
        ("Aman_123", "Aman_123"),
    ],
)
def test_sanitize_filename(raw, hasil):
    assert sanitize_filename(raw) == hasil


def test_template_tidak_ada(tmp_path):
    with pytest.raises(FileNotFoundError):
        generate_docx(tmp_path / "x.docx", tmp_path, "A", "1", {})


def test_nama_dan_folder_output_default(tmp_path):
    tpl = pilih_template_berdasarkan_pembimbing(1)
    out = generate_docx(tpl, tmp_path, "Andi/Wijaya", "21:12", {})
    assert out.parent == tmp_path / "AndiWijaya_2112"
    assert out.name == "Berita Acara dan Nilai Ujian Skripsi_AndiWijaya_2112.docx"
    assert out.exists()


def test_nama_file_custom(tmp_path):
    tpl = pilih_template_berdasarkan_pembimbing(1)
    out = generate_docx(tpl, tmp_path, "A", "1", {}, output_filename="Kustom/1")
    assert out.name == "Kustom1.docx"


@pytest.mark.parametrize("n", [1, 2])
def test_berita_acara_terisi(tmp_path, context_ba, n):
    ctx = dict(context_ba)
    if n == 1:
        ctx["pembimbing_2"] = ""
    out = generate_docx(
        pilih_template_berdasarkan_pembimbing(n),
        tmp_path,
        ctx["nama_mahasiswa"],
        ctx["npm"],
        ctx,
    )
    teks = teks_docx(out)
    assert "{{" not in teks, "masih ada placeholder yang belum terisi"
    for kunci in ("nama_mahasiswa", "npm", "judul_skripsi", "pembimbing_1",
                  "penguji_1", "penguji_2", "hari", "tanggal_bulan_tahun"):
        assert ctx[kunci] in teks, f"{kunci} tidak muncul di Berita Acara"
    assert ("Dr. Citra Dewi" in teks) == (n == 2)


@pytest.mark.parametrize("n", [1, 2])
def test_nota_dinas_terisi(tmp_path, context_nd, n):
    ctx = dict(context_nd)
    if n == 1:
        ctx["pembimbing_2"] = ""
    out = generate_docx(
        pilih_template_nota_dinas_berdasarkan_pembimbing(n),
        tmp_path,
        ctx["nama_mahasiswa"],
        ctx["npm"],
        ctx,
        output_filename="nd",
    )
    teks = teks_docx(out)
    assert "{{" not in teks
    for kunci in ("id_nd", "prodi", "lokasi_ujian", "tanggal_hari_ini",
                  "tahun_nd", "nama_mahasiswa", "npm", "judul_skripsi"):
        assert ctx[kunci] in teks, f"{kunci} tidak muncul di Nota Dinas"
